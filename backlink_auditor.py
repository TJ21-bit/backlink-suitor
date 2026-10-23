#!/usr/bin/env python3
"""
╔══════════════════════════════════════════════════════════╗
║   BACKLINK AUDITOR v4 — Hybrid Engine                    ║
║   Fast parallel HTTP first, headless browser only when   ║
║   the page actually needs it. No external services.      ║
╚══════════════════════════════════════════════════════════╝

HOW TO RUN:
    python backlink_auditor.py "Tijeet singh.xlsx"

FIRST TIME SETUP (run once):
    pip install -r requirements.txt
    python -m playwright install chromium

HOW IT WORKS (same flow as the Laravel CheckBacklinkJob):
    0. Pre-checks       — empty / invalid URL, self-link, file submissions
    1. HTTP phase       — 10 parallel plain HTTP requests (fast)
    2. Smart fallback   — a row goes to the headless browser only if:
                            • HTTP 401/403/404/429/5xx or connection timeout
                            • HTML too small / incomplete page
                            • CAPTCHA / Cloudflare challenge detected
                            • project domain not in a small page (JS-rendered)
                            • backlink or keyword not found in raw HTML
                            • known JS platform (Sngine, Pinterest, Reddit ...)
    3. Browser phase    — 3 parallel headless Chrome tabs, deep DOM scan
    4. Last resort      — alternate HTTP client (Safari UA, Bing referer)

CHECKS PER BACKLINK:
    • Live / dead / blocked / login wall, with exact reason
    • Soft 403 / 404 / 410 — error code but backlink is on the page
    • Project URL linked (domain, root domain, exact submitted URL, short links)
    • rel = dofollow / nofollow / sponsored / ugc, and anchor text
    • Target keyword (title, meta, anchor, H1-H3, body, JS-rendered text)
    • Content quality score (0-10) + AI-likelihood % with the signals found
"""

import argparse
import json
import math
import os
import queue
import re
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from urllib.parse import urlparse, urlsplit, urlunsplit, unquote

import pandas as pd
import requests
import urllib3
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

# On Render, keep Chrome inside the project folder (the default ~/.cache
# location is not kept between Render's build and run steps).
if os.environ.get("RENDER") and not os.environ.get("PLAYWRIGHT_BROWSERS_PATH"):
    os.environ["PLAYWRIGHT_BROWSERS_PATH"] = os.path.join(
        os.path.dirname(os.path.abspath(__file__)), ".playwright")

try:
    from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeout
    PLAYWRIGHT_OK = True
except ImportError:                                  # HTTP-only mode
    PLAYWRIGHT_OK = False

    class PlaywrightTimeout(Exception):
        pass

# Flush every print immediately so the web UI's live log keeps up
try:
    sys.stdout.reconfigure(line_buffering=True)
except Exception:
    pass


# ══════════════════════════════════════════════════════════════════════════════
# 0. CONFIG
# ══════════════════════════════════════════════════════════════════════════════

HTTP_WORKERS    = int(os.environ.get("HTTP_WORKERS", 10))      # parallel plain-HTTP requests
BROWSER_WORKERS = int(os.environ.get("BROWSER_WORKERS", 2 if os.environ.get("RENDER") else 3))  # parallel Chrome tabs
PAGE_TIMEOUT_MS = int(os.environ.get("PAGE_TIMEOUT", 20)) * 1000  # max wait for one page to load
MAX_BODY_BYTES  = 768 * 1024  # stop reading huge pages after this much

CLR_HEADER = "FF1565C0"
CLR_GREEN  = "FFB9F6CA"
CLR_RED    = "FFFFCDD2"
CLR_YELLOW = "FFFFF9C4"
CLR_BLUE   = "FFE3F2FD"
CLR_WHITE  = "FFFFFFFF"

RESULT_COLUMNS = [
    "Link Status", "Final URL", "Fetch Method", "Keyword Found?",
    "Project URL Linked?", "Link Rel", "Anchor Text",
    "Content Quality", "AI Content Score", "AI Signals",
    "Audit Result", "Flag Reason",
]

# Main client — looks like Chrome on Windows coming from Google
MAIN_HEADERS = {
    "User-Agent"               : "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
    "Accept"                   : "text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8",
    "Accept-Language"          : "en-US,en;q=0.9",
    "Accept-Encoding"          : "gzip, deflate",   # no 'br' — requests can't decode it without brotli
    "Connection"               : "keep-alive",
    "Referer"                  : "https://www.google.com/",
    "Cache-Control"            : "max-age=0",
    "Upgrade-Insecure-Requests": "1",
    "Sec-Fetch-Dest"           : "document",
    "Sec-Fetch-Mode"           : "navigate",
    "Sec-Fetch-Site"           : "cross-site",
    "Sec-Fetch-User"           : "?1",
    "sec-ch-ua"                : '"Chromium";v="124", "Google Chrome";v="124"',
    "sec-ch-ua-mobile"         : "?0",
    "sec-ch-ua-platform"       : '"Windows"',
}

# Alternate client — looks like Safari on Mac coming from Bing
ALT_HEADERS = {
    "User-Agent"     : "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/605.1.15 (KHTML, like Gecko) Version/17.0 Safari/605.1.15",
    "Accept"         : "text/html,application/xhtml+xml",
    "Accept-Language": "en-GB,en;q=0.9",
    "Accept-Encoding": "gzip, deflate",
    "Referer"        : "https://www.bing.com/",
}

HTTP_CODE_TEXT = {
    400: "Bad request — URL may be malformed",
    401: "Login required",
    403: "Forbidden — bot blocked",
    404: "Page not found",
    410: "Page permanently removed",
    429: "Too many requests — rate limited",
    500: "Server error",
    502: "Bad gateway",
    503: "Service unavailable",
    508: "Resource limit reached",
    509: "Bandwidth exceeded",
}


# ══════════════════════════════════════════════════════════════════════════════
# 1. DOMAIN LISTS
# ══════════════════════════════════════════════════════════════════════════════

SHORT_URL_DOMAINS = {
    "bit.ly", "tinyurl.com", "t.co", "t.ly", "goo.gl", "ow.ly", "is.gd",
    "cutt.ly", "rebrand.ly", "shorturl.at", "tiny.cc", "bl.ink", "buff.ly",
    "dlvr.it", "fur.ly", "mcaf.ee", "short.io", "soo.gd", "s.id", "v.gd",
    "clck.ru", "qr.ae", "adf.ly", "bit.do", "rb.gy",
    "pin.it", "flic.kr", "img.ly",
}

# Paste/document sites host content — never treat them as redirectors
PASTE_SITE_DOMAINS = {
    "justpaste.me", "justpaste.it", "pastebin.com", "paste.ee",
    "rentry.co", "privatebin.net", "hastebin.com", "dpaste.org",
}

# Sites that render their content with JavaScript — skip plain HTTP, go
# straight to the browser (from the Laravel job's list)
JS_PLATFORM_DOMAINS = {
    "pearltrees.com", "quora.com", "reddit.com", "pinterest.com", "pinterest.in",
    "flipboard.com", "scoop.it", "mix.com", "tumblr.com", "bizbuildboom.com",
    "diigo.com", "globalfreetalk.com", "streetvibex.com", "bondhusova.com",
    "instaconnect.co", "bookmarkingfree.com", "hasster.com", "logcla.com",
    "ooyy.com", "cherfanclub.com", "clarinetu.com", "claimajob.com",
    "livediscovery.com", "freewebmarks.com", "ybookmarking.com", "pikumil.com",
    "lalinguanostra.com", "yoursocial.it", "yoursocial.it.com", "zirvede.net",
    "woxmax.com", "freebookmarkingsite.com", "newsocialbookmarkingsite.com",
    "free-socialbookmarking.com", "social-bookmarkingsites.com",
    "freebookmarking.site", "bookmarking.win", "bookmark-dofollow.com",
    "vintfint.com",
}

# Sngine/PHPDolphin-style platforms that hide outbound links from visitors
# who are not logged in — an unlinked result there becomes "Manual Check"
LOW_QUALITY_PLATFORMS = {
    "fri3nd.me", "followingbook.com", "culturesbook.com", "topkif.nvinio.com",
    "we2chat.net", "antspride.com", "bytevidsocial.com", "instaconnect.co",
    "lebanonhub.app", "druzefaces.com", "yoursocial.it.com", "zirvede.net",
    "pikumil.com", "get2knowgod.com", "mn.co", "imo.net", "whoosmind.com",
    "ekcochat.com", "stayhomenetwork.com", "purekonect.com",
}

PDF_HOSTING_DOMAINS = {
    "dropbox.com", "4shared.com", "slideshare.net", "scribd.com", "issuu.com",
    "calameo.com", "docdroid.net", "mediafire.com", "docs.google.com",
    "drive.google.com", "heyzine.com", "files.fm", "smallpdf.com",
    "flipsnack.com", "yumpu.com", "pubhtml5.com", "joomag.com",
    "flipbuilder.com", "jumpshare.com", "filestack.com", "box.com",
    "onedrive.live.com",
}

IMAGE_HOSTING_DOMAINS = {
    "behance.net", "flickr.com", "pinterest.com", "instagram.com", "500px.com",
    "deviantart.com", "imgur.com", "photobucket.com", "unsplash.com", "pexels.com",
}

FILE_SUBMISSION_TYPES = {"pdf submission", "ppt submission", "image sharing", "video submission"}

# Bot-check / CAPTCHA fingerprints. Weak signals only count on near-empty
# pages; strong ones count anywhere.
CHALLENGE_SIGNALS = [
    "checking your browser", "ddos protection by cloudflare", "cf-browser-verification",
    "enable javascript and cookies", "just a moment", "cloudflare to restrict access",
    "__cf_chl", "cf_clearance", "recaptcha", "g-recaptcha", "hcaptcha", "grecaptcha",
    "your ip has been flagged", "security check to access", "verify you are human",
    "complete the security check", "automated access", "one more step",
    "attention required", "bot protection", "please wait", "stackprotectform",
    "security verification", "please wait while we verify", "verify you're not a robot",
    "verify you are not a robot", "we verify you", "spinner-border", "not a robot",
    "human verification", "access denied", "blocked",
]
STRONG_CHALLENGE_SIGNALS = {
    "recaptcha", "g-recaptcha", "hcaptcha", "grecaptcha",
    "__cf_chl", "cf_clearance", "stackprotectform",
}


# ══════════════════════════════════════════════════════════════════════════════
# 2. SMALL HELPERS
# ══════════════════════════════════════════════════════════════════════════════

def host_of(url: str) -> str:
    try:
        host = (urlparse(url).netloc or "").lower().split("@")[-1].split(":")[0]
        return host[4:] if host.startswith("www.") else host
    except Exception:
        return ""


def root_of(host: str) -> str:
    """blog.example.com → example.com (leaves co.uk-style hosts alone)."""
    parts = host.split(".")
    if len(parts) < 3:
        return host
    if parts[-2] in ("co", "com", "org", "net", "gov", "ac", "edu") and len(parts[-1]) == 2:
        return ".".join(parts[-3:])
    return ".".join(parts[-2:])


def in_domain_set(url: str, domain_set: set) -> bool:
    host = host_of(url)
    if not host:
        return False
    if host in domain_set:
        return True
    parts = host.split(".")
    return any(".".join(parts[k:]) in domain_set for k in range(1, len(parts) - 1))


def _is_short_url(url: str) -> bool:
    host = host_of(url)
    return host not in PASTE_SITE_DOMAINS and host in SHORT_URL_DOMAINS


def _is_pdf_url(url: str) -> bool:
    return in_domain_set(url, PDF_HOSTING_DOMAINS) or url.lower().split("?")[0].endswith(".pdf")


def _is_image_url(url: str) -> bool:
    return in_domain_set(url, IMAGE_HOSTING_DOMAINS)


def is_low_quality_platform(url: str) -> bool:
    return in_domain_set(url, LOW_QUALITY_PLATFORMS)


def is_js_platform(url: str) -> bool:
    return in_domain_set(url, JS_PLATFORM_DOMAINS)


def clean_url(url: str) -> str:
    """Trim and collapse double slashes in the path (https://site.com//post/x)."""
    url = (url or "").strip()
    try:
        p = urlsplit(url)
        return urlunsplit((p.scheme, p.netloc, re.sub(r"/{2,}", "/", p.path), p.query, p.fragment))
    except Exception:
        return url


_TAG_STRIP_RE = re.compile(r"<(script|style|noscript)\b.*?</\1\s*>", re.I | re.S)


def plain_text_fast(html: str) -> str:
    text = _TAG_STRIP_RE.sub(" ", html or "")
    text = re.sub(r"<[^>]+>", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def detect_challenge(html: str):
    """Return the CAPTCHA / bot-check signal found, or None."""
    if not html:
        return None
    lower = html.lower()
    m = re.search(r"<title[^>]*>(.*?)</title>", lower, re.S)
    title = m.group(1).strip() if m else ""
    for signal in ("just a moment", "attention required", "access denied", "security check",
                   "verify you are human", "human verification", "one more step", "bot protection"):
        if signal in title:
            return signal
    plain_size = len(plain_text_fast(html))
    for signal in CHALLENGE_SIGNALS:
        if signal in lower:
            if plain_size < 1000 or signal in STRONG_CHALLENGE_SIGNALS:
                return signal
    return None


def looks_like_challenge(html: str):
    """Stricter check: the page *is* a bot wall (not a normal page with a reCAPTCHA widget)."""
    signal = detect_challenge(html)
    if not signal:
        return None
    m = re.search(r"<title[^>]*>(.*?)</title>", (html or "").lower(), re.S)
    title = m.group(1).strip() if m else ""
    if signal in title or len(plain_text_fast(html)) < 1500:
        return signal
    return None


def is_login_page(html: str, url: str) -> bool:
    """True only when the page itself is a login wall (not just a login button)."""
    if not html:
        return False
    url_lower = (url or "").lower()
    for s in ("/login", "/signin", "/sign-in", "/auth/login", "/account/login", "/users/sign_in"):
        if url_lower.endswith(s) or f"{s}?" in url_lower:
            return True
    soup = BeautifulSoup(html, "html.parser")
    title = (soup.title.string or "").strip().lower() if soup.title and soup.title.string else ""
    if title in ("login", "sign in", "signin", "log in", "register", "create account", "join"):
        return True
    html_lower = html.lower()
    has_password = 'type="password"' in html_lower or "type='password'" in html_lower
    if not has_password:
        return False
    for tag in soup(["header", "footer", "nav", "script", "style", "aside"]):
        tag.decompose()
    return len(soup.get_text(separator=" ", strip=True)) < 500


# ── Short links ──────────────────────────────────────────────────────────────
# Resolved short links are saved to short_links.json next to this script, so
# each one is looked up once and reused on every later run. If a shortener
# blocks the tool, open the short link in your own browser and add a line:
#     "https://shorturl.at/AbCd1": "https://yoursite.com/the-page/"

SHORT_LINKS_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "short_links.json")
_SHORT_LOCK  = threading.Lock()


def _load_short_cache() -> dict:
    try:
        with open(SHORT_LINKS_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
        return {k.strip(): v.strip() for k, v in data.items() if isinstance(v, str) and v.strip()}
    except Exception:
        return {}


_SHORT_CACHE = _load_short_cache()


def _remember_short(url: str, dest: str):
    with _SHORT_LOCK:
        _SHORT_CACHE[url] = dest
        for k, v in _load_short_cache().items():     # keep entries other workers saved
            _SHORT_CACHE.setdefault(k, v)
        try:
            with open(SHORT_LINKS_FILE, "w", encoding="utf-8") as f:
                json.dump(dict(sorted(_SHORT_CACHE.items())), f, indent=2)
        except Exception:
            pass


def _short_key(url: str) -> str:
    return url.strip().rstrip("/")


def _resolve_short_url(url: str) -> str:
    """Follow a short link to its destination; returns the input if it can't."""
    key = _short_key(url)
    with _SHORT_LOCK:
        if key in _SHORT_CACHE:
            return _SHORT_CACHE[key]

    headers = {k: MAIN_HEADERS[k] for k in ("User-Agent", "Accept", "Accept-Language", "Accept-Encoding")}
    s = requests.Session()
    current = url
    # Hop by hop: read each Location header without loading the destination,
    # so a slow or bot-blocking destination site can't make the lookup fail.
    for _ in range(10):
        try:
            r = s.get(current, headers=headers, timeout=(4, 8), allow_redirects=False,
                      verify=False, stream=True)
        except Exception:
            break
        loc = r.headers.get("Location")
        if r.is_redirect and loc:
            r.close()
            current = requests.compat.urljoin(current, loc)
            if not _is_short_url(current) and host_of(current) != host_of(url):
                break                                      # left the shortener — done
            continue
        # No HTTP redirect — look for meta-refresh / JS redirect in the page
        try:
            body = r.raw.read(20000, decode_content=True).decode("utf-8", "replace")
        except Exception:
            body = ""
        r.close()
        m = (re.search(r'http-equiv=["\']?refresh["\']?[^>]*url=([^"\'>\s]+)', body, re.I)
             or re.search(r'(?:window\.)?location(?:\.href)?\s*=\s*["\'](https?://[^"\']+)', body, re.I)
             or re.search(r'<a[^>]+id=["\']?(?:destination|redirect|target)[^>]*href=["\'](https?://[^"\']+)', body, re.I))
        if m:
            current = requests.compat.urljoin(current, m.group(1))
        break

    if host_of(current) and host_of(current) != host_of(url):
        _remember_short(key, current)
        return current
    return url


def _resolve_short_in_browser(page, url: str) -> str:
    """Open a short link in a spare tab and record where it redirects to."""
    key = _short_key(url)
    with _SHORT_LOCK:
        if key in _SHORT_CACHE:
            return _SHORT_CACHE[key]
    short_host = host_of(url)
    seen = []

    def on_request(req):
        try:
            if req.is_navigation_request() and host_of(req.url) not in (short_host, "") \
                    and not _is_short_url(req.url):
                seen.append(req.url)
        except Exception:
            pass

    tab = None
    try:
        tab = page.context.new_page()
        tab.on("request", on_request)
        try:
            tab.goto(url, wait_until="commit", timeout=12000)
        except Exception:
            pass                       # the destination may fail — the redirect is what we need
        for _ in range(4):             # give interstitial / JS redirects up to ~4s
            if seen or host_of(tab.url) != short_host:
                break
            tab.wait_for_timeout(1000)
        dest = seen[0] if seen else tab.url
    except Exception:
        dest = seen[0] if seen else url
    finally:
        try:
            if tab:
                tab.close()
        except Exception:
            pass

    if dest and host_of(dest) and host_of(dest) != short_host:
        _remember_short(key, dest)
        return dest
    return url


_DNS_CACHE = {}


def domain_exists_publicly(host: str):
    """Ask Google public DNS. True = exists, False = NXDOMAIN, None = couldn't tell."""
    if host in _DNS_CACHE:
        return _DNS_CACHE[host]
    answer = None
    try:
        r = requests.get("https://dns.google/resolve", params={"name": host, "type": "A"}, timeout=6)
        data = r.json()
        if data.get("Status") == 0 and data.get("Answer"):
            answer = True
        elif data.get("Status") == 3:
            answer = False                      # NXDOMAIN — domain doesn't exist
        elif data.get("Status") == 2:
            answer = "servfail"                 # domain's name servers are down
    except Exception:
        pass
    _DNS_CACHE[host] = answer
    return answer


def dns_failure(url: str):
    """(state, error) for a name-resolution failure, split into dead vs local-network block."""
    host = host_of(url)
    exists = domain_exists_publicly(host)
    if exists is True:
        return "blocked", ("Domain did not resolve on your network, but it exists publicly — "
                           "ISP/DNS may be blocking it")
    if exists == "servfail":
        return "dead", "Domain's name servers are not responding — site is down"
    if exists is False:
        return "dead", "Domain does not exist (expired or wrong URL)"
    return "dead", "Domain not found — URL may be wrong"


# ══════════════════════════════════════════════════════════════════════════════
# 3. FETCHERS — every fetcher returns the same shape:
#    { state: live|dead|blocked|login, code, html, final_url, error, method }
# ══════════════════════════════════════════════════════════════════════════════

def _decode(body: bytes, content_type: str) -> str:
    m = re.search(r"charset=([\w-]+)", content_type or "", re.I)
    enc = m.group(1) if m else None
    if not enc:
        m = re.search(rb'<meta[^>]+charset=["\']?([\w-]+)', body[:4096], re.I)
        enc = m.group(1).decode("ascii", "ignore") if m else "utf-8"
    try:
        return body.decode(enc, errors="replace")
    except LookupError:
        return body.decode("utf-8", errors="replace")


def http_fetch(url: str, alt: bool = False) -> dict:
    res = {"state": "blocked", "code": None, "html": "", "final_url": url,
           "error": None, "method": "alt-http" if alt else "http"}
    try:
        s = requests.Session()
        s.max_redirects = 8
        with s.get(url, headers=ALT_HEADERS if alt else MAIN_HEADERS,
                   timeout=(4, 8) if alt else (6, 12),
                   allow_redirects=True, verify=False, stream=True) as r:
            res["code"]      = r.status_code
            res["final_url"] = r.url or url
            buf = bytearray()
            deadline = time.time() + (15 if alt else 25)   # a server trickling bytes can't stall us
            for chunk in r.iter_content(8192):
                buf.extend(chunk)
                if len(buf) >= MAX_BODY_BYTES or time.time() > deadline:
                    break
            res["html"] = _decode(bytes(buf), r.headers.get("Content-Type", ""))
    except requests.exceptions.TooManyRedirects:
        res["error"] = "Too many redirects"
        return res
    except requests.exceptions.Timeout:
        res["error"] = "Connection timeout — site too slow or blocking bots"
        return res
    except requests.exceptions.ConnectionError as e:
        msg = str(e)
        if any(k in msg for k in ("Name or service not known", "getaddrinfo", "nodename nor servname",
                                  "NameResolutionError", "No address associated")):
            res["state"], res["error"] = dns_failure(url)
        elif "Connection refused" in msg:
            res["state"], res["error"] = "dead", "Connection refused"
        else:
            res["error"] = "Connection closed — server may be overloaded"
        return res
    except Exception as e:
        res["error"] = str(e)[:120]
        return res

    code = res["code"]
    if 200 <= code < 300:
        res["state"] = "live"
    elif code in (400, 404, 410):
        res["state"] = "dead"
        res["error"] = f"HTTP {code} — {HTTP_CODE_TEXT[code]}"
    else:
        res["state"] = "blocked"
        res["error"] = f"HTTP {code} — {HTTP_CODE_TEXT.get(code, 'unexpected response')}"
    return res


def needs_browser_reason(fetch: dict, project_domain: str):
    """The Laravel job's 'smart Puppeteer fallback' rules. None = HTTP is enough."""
    code = fetch["code"]
    html = fetch["html"] or ""
    if code is None:
        return None if fetch["state"] == "dead" else (fetch["error"] or "connection error")
    if code == 410:
        return None                                   # soft-410 is decided in analysis
    if not (200 <= code < 300):
        return f"HTTP {code}"
    if len(html.strip()) < 1000:
        return f"HTML too small ({len(html)} bytes)"
    challenge = detect_challenge(html)
    if challenge:
        return f"challenge page: {challenge}"
    lower = html.lower()
    plain = plain_text_fast(html)
    if ("</body>" not in lower and "</html>" not in lower) or len(plain) < 200:
        return "incomplete page content"
    if project_domain and project_domain not in lower and len(plain) < 15000:
        return "JS likely — project domain not in HTML"
    return None


# ── Browser fetch (Playwright) ───────────────────────────────────────────────

def _scroll_page(page):
    try:
        page.evaluate("""
            async () => {
                window.scrollTo(0, document.body.scrollHeight / 2);
                await new Promise(r => setTimeout(r, 300));
                window.scrollTo(0, document.body.scrollHeight);
                await new Promise(r => setTimeout(r, 200));
                window.scrollTo(0, 0);
            }
        """)
        page.wait_for_timeout(200)
    except Exception:
        pass


def _grab_html(page) -> str:
    html = page.content()
    for frame in page.frames[1:]:
        try:
            html += frame.content()
        except Exception:
            pass
    return html


def browser_fetch(page, url: str) -> dict:
    res = {"state": "blocked", "code": None, "html": "", "final_url": url,
           "error": None, "method": "browser"}
    actual = url
    if _is_short_url(url):
        resolved = _resolve_short_url(url)
        if resolved != url:
            actual = res["final_url"] = resolved

    def finish_live(code):
        _scroll_page(page)
        html = _grab_html(page)
        res["code"], res["final_url"] = code, page.url or actual
        if is_login_page(html, page.url):
            res["state"] = "login"
            res["error"] = "Login required — content hidden behind login wall"
        else:
            res["state"], res["html"] = "live", html

    try:
        resp = page.goto(actual, wait_until="domcontentloaded", timeout=PAGE_TIMEOUT_MS)
        code = resp.status if resp else None
        try:
            page.wait_for_load_state("networkidle", timeout=2500)
        except PlaywrightTimeout:
            pass

        if code is not None and 200 <= code < 300:
            finish_live(code)
        elif code is None:
            html = _grab_html(page)
            if len(html) > 1000:
                finish_live(200)
            else:
                res["error"] = "No HTTP status — page may have redirected"
        elif code in (429, 500, 502, 503, 508, 509):
            time.sleep(2)                               # temporary — retry once
            resp2 = page.goto(actual, wait_until="domcontentloaded", timeout=PAGE_TIMEOUT_MS)
            if resp2 and resp2.status == 200:
                finish_live(200)
            else:
                res["code"] = code
                res["error"] = f"HTTP {code} — {HTTP_CODE_TEXT.get(code, 'server error')}"
        else:
            res["code"] = code
            res["html"] = _grab_html(page)              # kept for soft-404 checks
            res["state"] = "dead" if code in (400, 404, 410) else "blocked"
            res["error"] = f"HTTP {code} — {HTTP_CODE_TEXT.get(code, 'unexpected response')}"

    except PlaywrightTimeout:
        try:                                            # last try: first byte only
            resp = page.goto(actual, wait_until="commit", timeout=8000)
            page.wait_for_timeout(2000)
            html = _grab_html(page)
            code = resp.status if resp else None
            if len(html) > 1500 and (code is None or 200 <= code < 300):
                res.update(state="live", code=code or 200, html=html, final_url=page.url)
            else:
                res["error"] = "Timeout — site too slow or bot-blocking"
        except Exception:
            res["error"] = "Timeout — site too slow or bot-blocking"

    except Exception as e:
        err = str(e)
        if "ERR_NAME_NOT_RESOLVED" in err:
            res["state"], res["error"] = dns_failure(url)
        elif "ERR_CONNECTION_REFUSED" in err:
            res["state"], res["error"] = "dead", "Connection refused"
        elif "ERR_SSL" in err or "ERR_CERT" in err:
            res["error"] = "SSL certificate error"
        elif "TIMED_OUT" in err:
            res["error"] = "Connection timed out — server unreachable"
        elif "ERR_CONNECTION_CLOSED" in err or "ERR_CONNECTION_RESET" in err:
            res["error"] = "Connection closed — server may be overloaded"
        else:
            res["error"] = err.replace("Page.goto: ", "").split(" at http")[0][:100] or "Unknown error"

    return res


# ══════════════════════════════════════════════════════════════════════════════
# 4. KEYWORD CHECK
# ══════════════════════════════════════════════════════════════════════════════

def check_keyword(html: str, keyword: str) -> tuple:
    """(found, where). Also tries 'skil travel' / 'skil-travel' / 'skiltravel'."""
    kw = keyword.strip().lower()
    if not kw or not html:
        return False, None
    soup = BeautifulSoup(html, "html.parser")

    variants = [kw]
    if " " in kw:
        variants += [kw.replace(" ", ""), kw.replace(" ", "-")]
    else:
        for i in range(3, len(kw) - 2):
            variants += [kw[:i] + " " + kw[i:], kw[:i] + "-" + kw[i:]]
    variants = list(dict.fromkeys(variants))

    title_text = (soup.title.string or "").lower() if soup.title and soup.title.string else ""
    meta_tag   = soup.find("meta", attrs={"name": "description"})
    meta_text  = (meta_tag.get("content") or "").lower() if meta_tag else ""
    anchors    = [(a.get_text(strip=True) or "").lower() for a in soup.find_all("a")]
    headings   = [(h.get_text(strip=True).lower(), h.name.upper()) for h in soup.find_all(["h1", "h2", "h3"])]
    body_text  = soup.get_text(separator=" ", strip=True).lower()

    for v in variants:
        if v in title_text:
            return True, "Page Title"
        if v in meta_text:
            return True, "Meta Description"
        if any(v in a for a in anchors):
            return True, "Anchor Text"
        for text, tag in headings:
            if v in text:
                return True, tag
        if v in body_text:
            return True, "Page Body"
    return False, None


# ══════════════════════════════════════════════════════════════════════════════
# 5. PROJECT URL LINK CHECK
# ══════════════════════════════════════════════════════════════════════════════

def _rel_info(rel_values) -> dict:
    rel = {v.lower() for v in (rel_values or [])}
    return {
        "nofollow" : "nofollow" in rel,
        "sponsored": "sponsored" in rel,
        "ugc"      : "ugc" in rel,
        "dofollow" : not ({"nofollow", "sponsored"} & rel),
        "known"    : True,
    }


UNKNOWN_REL = {"nofollow": False, "sponsored": False, "ugc": False, "dofollow": True, "known": False}


def check_project_url_linked(html: str, project_url: str) -> dict:
    """
    Does the page link to the project? Returns
    { linked, rel, link_text, match_type, exact }
      rel   — dofollow / nofollow / sponsored / ugc info for the matched <a>
      exact — the submitted URL's path (not just the domain) is on the page
    """
    empty = {"linked": False, "rel": UNKNOWN_REL, "link_text": "—", "match_type": None, "exact": False}
    if not project_url or not html:
        return empty

    soup           = BeautifulSoup(html, "html.parser")
    project_domain = host_of(project_url)
    project_root   = root_of(project_domain)
    project_path   = urlparse(project_url).path.rstrip("/").lower()
    project_full   = project_url.rstrip("/").lower()
    project_bare   = re.sub(r"^https?://(www\.)?", "", project_full)
    if not project_domain:
        return empty

    html_decoded = unquote(html.lower())
    hrefs        = [unquote((a.get("href") or "").lower()) for a in soup.find_all("a", href=True)]
    if project_path and len(project_path) > 1:
        exact = project_bare in html_decoded or any(project_path in h and project_domain in h for h in hrefs)
    else:
        exact = True                                    # homepage — domain is the URL

    def hit(rel, text, match):
        return {"linked": True, "rel": rel, "link_text": (text or "—")[:80],
                "match_type": match, "exact": exact}

    unresolved_shorts = []        # short links we couldn't follow over plain HTTP

    def short_points_to_project(short):
        dest = _resolve_short_url(short)
        if host_of(dest) == host_of(short):
            if short not in unresolved_shorts:
                unresolved_shorts.append(short)
            return False
        return project_domain in dest.lower()

    # Stage 1 — real <a href> links
    for a in soup.find_all("a", href=True):
        raw  = a.get("href") or ""
        href = unquote(raw.lower())
        h    = host_of(href)
        match = None
        if h == project_domain:
            match = "Direct link"
        elif project_root != project_domain and (h == project_root or h.endswith("." + project_root)):
            match = "Root domain link"
        elif project_bare and project_bare in href:
            match = "Via redirect"
        elif project_domain in href:
            match = "Domain in link"
        elif _is_short_url(raw):
            if short_points_to_project(raw.strip()):
                match = "Short link (resolved)"
        if match:
            return hit(_rel_info(a.get("rel")), a.get_text(strip=True), match)

    # Stage 2 — other attributes that carry URLs (iframes, data-href ...)
    for tag in soup.find_all(True):
        for attr in ("data-href", "data-url", "data-link", "data-original-url",
                     "data-target", "src", "data-src"):
            val = tag.get(attr)
            if isinstance(val, str) and project_domain in unquote(val.lower()):
                return hit(UNKNOWN_REL, tag.get_text(strip=True) or tag.get("title", ""),
                           f"Via {tag.name} {attr}")

    # Stage 3 — short URLs written as plain text ("Read more: https://bit.ly/x")
    page_text = soup.get_text(separator=" ", strip=True)
    for found in re.findall(r'https?://[^\s<>"\')\]]+', page_text):
        found = found.rstrip(".,;:!?)")
        if _is_short_url(found) and short_points_to_project(found):
            return hit(UNKNOWN_REL, f"Short URL in text → {found}", "Short URL in text (resolved)")

    # Stages 4-5 only count occurrences that are in the markup/scripts, not
    # mentions a reader sees as plain text (those are not backlinks).
    text_lower = unquote(page_text.lower())

    def in_markup(needle):
        return bool(needle) and html_decoded.count(needle) > text_lower.count(needle)

    # Stage 4 — full submitted URL embedded in the source (JS-built links)
    if in_markup(project_bare):
        return hit(UNKNOWN_REL, "Found in page source", "JS / embedded")

    # Stage 5 — domain in the source next to link-ish markup
    if in_markup(project_domain):
        for m in re.finditer(re.escape(project_domain), html_decoded):
            context = html_decoded[max(0, m.start() - 200): m.start() + 200]
            if any(sig in context for sig in ("href", "src=", "url", "http")):
                return hit(UNKNOWN_REL, "Domain found in page source", "Domain in source")

    # Short links we couldn't follow — the backlink may well be behind them
    if unresolved_shorts:
        return {"linked": False, "rel": UNKNOWN_REL, "link_text": ", ".join(unresolved_shorts[:3]),
                "match_type": "short-unresolved", "exact": exact, "short_urls": unresolved_shorts}

    # Stage 6 — only mentioned as plain text: not a backlink
    if project_domain in page_text.lower():
        return {"linked": False, "rel": UNKNOWN_REL, "link_text": "URL mentioned but NOT clickable",
                "match_type": "text-only", "exact": exact}

    return empty


def _deep_retry_project_url(page, project_url: str) -> dict:
    """Post-render DOM scan for links injected by AJAX after load."""
    empty = {"linked": False, "rel": UNKNOWN_REL, "link_text": "—", "match_type": None, "exact": False}
    try:
        page.wait_for_timeout(800)
        page.evaluate("""
            async () => {
                const sleep = ms => new Promise(r => setTimeout(r, ms));
                for (const f of [0, 0.5, 1, 0.5, 0]) {
                    window.scrollTo(0, document.body.scrollHeight * f); await sleep(200);
                }
            }
        """)
        page.wait_for_timeout(600)

        project_domain = host_of(project_url)
        project_path   = urlparse(project_url).path.rstrip("/").lower()
        anchors = page.evaluate("""
            () => Array.from(document.querySelectorAll('a')).map(a => ({
                href: a.href || '',
                text: (a.innerText || a.textContent || '').trim().slice(0, 80),
                rel : a.getAttribute('rel') || ''
            }))
        """)
        for a in anchors:
            href = unquote((a.get("href") or "").lower())
            if host_of(href) == project_domain or project_domain in href:
                exact = (not project_path or len(project_path) <= 1) or project_path in href
                return {"linked": True, "rel": _rel_info((a.get("rel") or "").split()),
                        "link_text": (a.get("text") or "—")[:80], "match_type": "Deep scan (rendered DOM)",
                        "exact": exact}

        if project_domain in unquote(page.content().lower()):
            return {"linked": True, "rel": UNKNOWN_REL, "link_text": "Found in rendered page",
                    "match_type": "Deep scan (rendered source)", "exact": False}
    except Exception:
        pass
    return empty


# ══════════════════════════════════════════════════════════════════════════════
# 6. CONTENT QUALITY + AI LIKELIHOOD (ported from the Laravel scoreContent)
# ══════════════════════════════════════════════════════════════════════════════

AI_PHRASES = [
    ("delve into", 3), ("delves into", 3), ("delving into", 3),
    ("in today's", 3), ("in the realm of", 3), ("in the world of", 3),
    ("in the landscape of", 3), ("navigate the complex", 3),
    ("multifaceted", 3), ("pivotal role", 3), ("plays a pivotal", 3),
    ("plays a crucial", 3), ("plays a vital", 3), ("plays a key role", 3),
    ("it is important to note", 3), ("it is worth noting", 3),
    ("a testament to", 3), ("to new heights", 3), ("elevate your", 3),
    ("harness the power", 3), ("tapestry", 3), ("rich tapestry", 3),
    ("ever-evolving", 3), ("rapidly evolving", 3),
    ("embark on a journey", 3), ("take a deep dive", 3),
    ("treasure trove", 3), ("unlock the potential", 3),
    ("at the forefront", 3), ("game-changer", 3), ("game-changing", 3),
    ("cutting-edge", 3), ("paradigm shift", 3), ("revolutionize", 3),
    ("unparalleled", 3), ("meticulously", 3), ("meticulously crafted", 3),
    ("seamlessly integrate", 3), ("a plethora of", 3), ("a myriad of", 3),
    ("it goes without saying", 3), ("a wide array of", 3),
    ("by leveraging", 3), ("by harnessing", 3),
    ("look no further", 3), ("let's dive into", 3),
    ("in this blog post", 3), ("dynamic landscape", 3),
    ("digital landscape", 3), ("synergy", 3), ("synergies", 3),
    ("furthermore", 2), ("moreover", 2), ("additionally", 2),
    ("consequently", 2), ("in conclusion", 2), ("to sum up", 2),
    ("in summary", 2), ("ultimately", 2), ("leverage", 2),
    ("robust", 2), ("seamless", 2), ("holistic", 2), ("comprehensive", 2),
    ("paramount", 2), ("pivotal", 2), ("undoubtedly", 2),
    ("remarkable", 2), ("noteworthy", 2), ("intricate", 2),
    ("foster", 2), ("cultivate", 2), ("streamline", 2),
    ("empower", 2), ("facilitate", 2), ("transformative", 2),
    ("bespoke", 2), ("curated", 2), ("evolving landscape", 2),
    ("it is essential", 2), ("one of the key", 2), ("one of the most", 2),
    ("when it comes to", 2), ("whether you are", 2),
    ("we will explore", 2), ("navigating the", 2),
    ("state-of-the-art", 2), ("unprecedented", 2),
    ("crucial", 1), ("vital", 1), ("enhance", 1),
    ("optimize", 1), ("tailored", 1),
]
# Word-start boundary: 'foster' matches 'fostering' but 'thus' won't match 'enthusiast'
_AI_PHRASE_RES = [(p, w, re.compile(r"\b" + re.escape(p))) for p, w in AI_PHRASES]

AI_FINGERPRINT_WORDS = {
    "delve", "delves", "delved", "delving", "tapestry", "tapestries", "underscore", "underscores",
    "elevate", "elevates", "elevated", "elevating", "leverage", "leverages", "leveraged", "leveraging",
    "harness", "harnesses", "harnessed", "harnessing", "embark", "embarks", "embarked", "embarking",
    "navigate", "navigates", "navigated", "navigating", "foster", "fosters", "fostered", "fostering",
    "cultivate", "cultivates", "cultivated", "cultivating", "streamline", "streamlines", "streamlined",
    "streamlining", "empower", "empowers", "empowered", "empowering", "facilitate", "facilitates",
    "facilitated", "facilitating", "encompass", "encompasses", "encompassed", "encompassing",
    "myriad", "plethora", "multitude", "paramount", "pivotal", "robust", "seamless", "holistic",
    "transformative", "groundbreaking", "revolutionary", "intricate", "multifaceted", "nuanced",
    "remarkable", "noteworthy", "meticulous", "meticulously", "synergy", "synergies", "synergistic",
    "bolster", "bolsters", "endeavor", "endeavors", "imperative", "indispensable", "invaluable",
    "ubiquitous", "showcase", "showcases", "spearhead", "epitome", "quintessential", "moreover",
    "furthermore", "additionally", "essentially", "fundamentally", "ultimately", "undoubtedly",
}

TRANSITION_WORDS = [
    "furthermore", "moreover", "additionally", "consequently", "however", "therefore",
    "thus", "hence", "in addition", "in contrast", "nonetheless", "nevertheless",
    "subsequently", "accordingly",
]
_TRANSITION_RES = [re.compile(r"\b" + re.escape(t) + r"\b") for t in TRANSITION_WORDS]
_TRIPLET_RE = re.compile(
    r"\b\w+(?:ly|ing|ed|tion|ness|ity)?,\s+\b\w+(?:ly|ing|ed|tion|ness|ity)?,?\s+and\s+\b\w+", re.I)


def _extract_page_text(html: str) -> str:
    """Body text without scripts, menus, headers and footers."""
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header", "aside", "form", "noscript"]):
        tag.decompose()
    return " ".join(soup.get_text(separator=" ", strip=True).split())


def _count_syllables(word: str) -> int:
    word = word.lower()
    if len(word) <= 3:
        return 1
    word = re.sub(r"(?:[^laeiouy]es|ed|[^laeiouy]e)$", "", word)
    word = re.sub(r"^y", "", word)
    return max(1, len(re.findall(r"[aeiouy]{1,2}", word)))


def score_content(text: str):
    """
    Returns { score, quality_label, ai_pct, origin_label, signals } or None
    if there isn't enough text. score = writing quality 0-10 minus an AI penalty.
    """
    text = (text or "").replace("’", "'").replace("‘", "'").strip()
    tokens = text.split(" ")
    if len(tokens) > 5000:
        text = " ".join(tokens[:5000])
    lower = text.lower()

    words = re.findall(r"\b[a-z']+\b", text, re.I)
    wc = len(words)
    if wc < 20:
        return None

    sentences = [s for s in re.split(r"[.!?]+\s+", text) if s.strip()]
    sc        = len(sentences)
    sent_lens = [len([w for w in s.split(" ") if w]) for s in sentences]
    avg_len   = sum(sent_lens) / sc if sc else 0
    sent_std  = math.sqrt(sum((l - avg_len) ** 2 for l in sent_lens) / sc) if sc > 1 else 0

    lower_words = [w.lower() for w in words]
    ttr = len(set(lower_words)) / wc
    syllables = sum(_count_syllables(w) for w in words)
    flesch = (206.835 - 1.015 * (wc / sc) - 84.6 * (syllables / wc)) if sc else 0

    bigrams = [lower_words[k] + " " + lower_words[k + 1] for k in range(wc - 1)]
    bigram_repeat = 1 - len(set(bigrams)) / len(bigrams) if bigrams else 0

    personal_density    = len(re.findall(r"\b(?:i|me|my|we|us|our|you|your)\b", text, re.I)) / wc
    contraction_density = len(re.findall(r"\b\w+'(?:s|t|re|ve|ll|d|m)\b", text, re.I)) / wc

    burstiness = 0
    if len(sent_lens) > 3:
        burstiness = sum(abs(sent_lens[k] - sent_lens[k - 1]) for k in range(1, len(sent_lens))) / (len(sent_lens) - 1)

    em_dash_count   = text.count("—")
    em_dash_density = em_dash_count / max(1, sc)

    phrase_hits, phrase_weight = [], 0
    for phrase, weight, rx in _AI_PHRASE_RES:
        n = len(rx.findall(lower))
        if n:
            phrase_hits.append((phrase, n, weight))
            phrase_weight += n * weight
    phrase_density = phrase_weight / max(1, wc / 100)

    fingerprint_hits    = [w for w in lower_words if w in AI_FINGERPRINT_WORDS]
    fingerprint_density = len(fingerprint_hits) / wc * 1000

    transition_density = sum(len(rx.findall(lower)) for rx in _TRANSITION_RES) / max(1, sc)
    triplet_density    = len(_TRIPLET_RE.findall(text)) / max(1, sc)

    # ── Writing quality (0-10) ──────────────────────────────────────────────
    def band(value, steps, default):
        for threshold, pts in steps:
            if value >= threshold:
                return pts
        return default

    len_score  = band(wc, [(600, 2), (300, 1.5), (150, 1), (50, 0.5)], 0)
    ttr_score  = band(ttr, [(0.55, 2), (0.45, 1.5), (0.35, 1), (0.25, 0.5)], 0)
    if 50 <= flesch <= 75:
        read_score = 2
    elif 35 <= flesch <= 85:
        read_score = 1.5
    elif 20 <= flesch <= 95:
        read_score = 1
    else:
        read_score = 0.5
    var_score  = band(sent_std, [(6, 2), (4, 1.5), (2.5, 1)], 0.5)
    rep_score  = 2 if bigram_repeat <= 0.05 else 1.5 if bigram_repeat <= 0.10 else 1 if bigram_repeat <= 0.20 else 0.5
    quality    = min(10, len_score + ttr_score + read_score + var_score + rep_score)
    q_details  = [
        f"Words:{wc}({len_score:.1f}/2)", f"Vocab:{ttr*100:.1f}%({ttr_score:.1f}/2)",
        f"Flesch:{flesch:.1f}({read_score:.1f}/2)", f"SentVariety:σ={sent_std:.1f}({var_score:.1f}/2)",
        f"Repeat:{bigram_repeat*100:.1f}%({rep_score:.1f}/2)",
    ]

    # ── AI likelihood ───────────────────────────────────────────────────────
    ai, signals = 0.0, []
    ai += min(45, phrase_density * 8)
    if phrase_hits:
        phrase_hits.sort(key=lambda h: -(h[1] * h[2]))
        top = ", ".join(f'"{p}"' + (f"(x{n})" if n > 1 else "") for p, n, _ in phrase_hits[:4])
        signals.append(f"{len(phrase_hits)} AI phrases: {top}")
    if sc > 4:
        if burstiness < 4:
            ai += 20; signals.append(f"Uniform rhythm(burst={burstiness:.1f})")
        elif burstiness < 6:
            ai += 12; signals.append(f"Low rhythm(burst={burstiness:.1f})")
        elif burstiness < 8:
            ai += 5
    if sc > 5:
        if sent_std < 3:
            ai += 12; signals.append(f"Low sentence variance(σ={sent_std:.1f})")
        elif sent_std < 5:
            ai += 6
    if transition_density > 0.4:
        ai += 15; signals.append(f"Excessive transitions({transition_density*100:.0f}%)")
    elif transition_density > 0.2:
        ai += 8; signals.append("Heavy transitions")
    elif transition_density > 0.1:
        ai += 3
    if triplet_density > 0.25:
        ai += 10; signals.append("Parallel triplets")
    elif triplet_density > 0.12:
        ai += 5
    if em_dash_density > 0.5:
        ai += 8; signals.append(f"Em-dash overuse({em_dash_count})")
    elif em_dash_density > 0.25:
        ai += 4
    if wc > 80:
        if personal_density < 0.005:
            ai += 8; signals.append("No personal pronouns")
        elif personal_density < 0.015:
            ai += 4
        if contraction_density < 0.003:
            ai += 8; signals.append("No contractions")
        elif contraction_density < 0.008:
            ai += 3
    if avg_len > 28:
        ai += 6; signals.append(f"Long sentences({avg_len:.1f}w)")
    elif avg_len > 22:
        ai += 3
    if fingerprint_hits:
        ai += min(35, fingerprint_density * 4)
        uniq = list(dict.fromkeys(fingerprint_hits))[:5]
        signals.append(f"{len(fingerprint_hits)} fingerprint words: {','.join(uniq)}")

    ai_pct = max(0.0, min(100.0, 100 / (1 + math.exp(-0.08 * (ai - 35)))))
    penalty = 4 if ai_pct >= 80 else 3 if ai_pct >= 65 else 2 if ai_pct >= 50 else 1 if ai_pct >= 35 else 0.5 if ai_pct >= 20 else 0
    score = round(max(0, min(10, quality - penalty)), 1)
    ai_pct = round(ai_pct)

    quality_label = ("Excellent Quality" if score >= 7.5 else "Good Quality" if score >= 6
                     else "Average Quality" if score >= 4 else "Low Quality")
    origin_label  = ("AI Generated" if ai_pct >= 75 else "Likely AI" if ai_pct >= 50
                     else "Mixed/Uncertain" if ai_pct >= 30 else "Likely Human" if ai_pct >= 15
                     else "Human Written")
    return {
        "score": score, "quality_label": quality_label, "ai_pct": ai_pct,
        "origin_label": origin_label,
        "signals": " | ".join(signals[:4] + q_details)[:500],
    }


# ══════════════════════════════════════════════════════════════════════════════
# 7. ANALYSE ONE FETCHED PAGE
# ══════════════════════════════════════════════════════════════════════════════

def _row_fields(row: dict):
    return (
        clean_url(str(row.get("Backlink Live Link", "") or "")),
        str(row.get("Target Keyword", "") or "").strip(),
        clean_url(str(row.get("Submitted Project Url", "") or "")),
        str(row.get("Backlink Activity Type", "") or "").strip(),
    )


def make_result(link_status, audit_result, reason, method="—", **overrides) -> dict:
    res = {
        "Link Status"        : link_status,
        "Final URL"          : "",
        "Fetch Method"       : method,
        "Keyword Found?"     : "⛔ N/A",
        "Project URL Linked?": "⛔ N/A",
        "Link Rel"           : "",
        "Anchor Text"        : "",
        "Content Quality"    : "⛔ N/A",
        "AI Content Score"   : "⛔ N/A",
        "AI Signals"         : "",
        "Audit Result"       : audit_result,
        "Flag Reason"        : reason,
    }
    res.update(overrides)
    return res


METHOD_LABELS = {"http": "HTTP", "alt-http": "HTTP (alt client)", "browser": "Headless browser"}


def analyze(row: dict, fetch: dict, page=None, note: str = "") -> dict:
    live_url, keyword, project_url, activity = _row_fields(row)
    is_onsite = activity.lower() == "on-site blog"
    flags, reasons = [], []
    manual = []                     # things a person must confirm (→ Manual Check, not Flagged)

    state     = fetch["state"]
    html      = fetch.get("html") or ""
    final_url = fetch.get("final_url") or live_url
    method    = METHOD_LABELS.get(fetch["method"], fetch["method"]) + (f" — {note}" if note else "")
    soft_note = None
    proj      = None

    # Soft 404/410: server sends an error code but the backlink is on the page
    if state == "dead" and html and project_url and not is_onsite:
        proj = check_project_url_linked(html, project_url)
        if proj["linked"]:
            state, soft_note = "live", f"Soft {fetch['code']} — server returns an error code but the backlink is on the page"
        else:
            proj = None
    if state == "live" and is_login_page(html, final_url):
        state = "login"

    redirected = host_of(final_url) and host_of(final_url) != host_of(live_url)
    final_url_out = final_url if redirected else ""

    # ── Link status ─────────────────────────────────────────────────────────
    if state == "live":
        link_status = "✅ Live"
        if redirected:
            link_status += f" (redirect → {host_of(final_url)})"
        if soft_note:
            link_status += f" — Soft {fetch['code']}"
            reasons.append(soft_note)
    elif state == "login":
        link_status = "⚠️ Login required — content hidden behind login wall"
        flags.append("Login Wall")
        reasons.append("Page shows a login screen — backlink not publicly visible, no SEO value")
    elif state == "blocked":
        link_status = f"⚠️ Verify manually — {fetch.get('error') or 'bot blocked'}"
    else:
        link_status = f"❌ {fetch.get('error') or 'Dead'}"
        flags.append("Dead Link")
        reasons.append(f"Dead link — {fetch.get('error') or 'page not reachable'}")

    na = {"dead": "⛔ N/A (link dead)", "login": "⛔ N/A (login required)",
          "blocked": "⚠️ Verify manually (blocked)"}.get(state)

    kw_display = proj_display = na or ""
    rel_display = anchor_display = ai_signals = ""
    quality_display = ai_display = na or ""
    kw_found = True
    proj_needs_retry = False
    content = None

    if state == "live":
        # ── Keyword ─────────────────────────────────────────────────────────
        if not keyword:
            kw_display = "⛔ N/A (no keyword)"
        elif _is_pdf_url(final_url):
            kw_display = "⛔ N/A (PDF — cannot check keyword in document)"
        else:
            kw_found, kw_where = check_keyword(html, keyword)
            if not kw_found and page is not None:
                try:
                    if keyword.lower() in page.inner_text("body").lower():
                        kw_found, kw_where = True, "Page Body (JS rendered)"
                except Exception:
                    pass
            if kw_found:
                kw_display = f"✅ Yes ({kw_where})"
            elif _is_image_url(final_url):
                kw_display = "⚠️ Image platform — keyword only in image title/description"
            elif is_low_quality_platform(final_url):
                kw_display = "⚠️ Verify manually (content may be behind login)"
            else:
                kw_display = "❌ Not found"
                flags.append("Keyword Missing")
                reasons.append(f'"{keyword}" not found on page')

        # ── Project URL ─────────────────────────────────────────────────────
        if is_onsite:
            proj_display = "⛔ N/A (on-site blog — internal page)"
        elif not project_url:
            proj_display = "⛔ N/A (no project URL)"
        else:
            if proj is None:
                proj = check_project_url_linked(html, project_url)
            if not proj["linked"] and proj["match_type"] != "text-only" and page is not None:
                deep = _deep_retry_project_url(page, project_url)
                if deep["linked"]:
                    proj = deep

            # Short links plain HTTP couldn't follow — open them in the browser
            if not proj["linked"] and proj["match_type"] == "short-unresolved" and page is not None:
                pdomain = host_of(project_url)
                still = []
                for short in proj.get("short_urls", [])[:5]:
                    dest = _resolve_short_in_browser(page, short)
                    if pdomain in dest.lower():
                        exact = urlparse(project_url).path.rstrip("/").lower() in dest.lower()
                        proj = {"linked": True, "rel": UNKNOWN_REL, "link_text": f"{short} → {host_of(dest)}",
                                "match_type": "Short link (resolved in browser)", "exact": exact}
                        break
                    if host_of(dest) == host_of(short):
                        still.append(short)
                else:
                    if not still:          # every short link resolved, none to the project
                        proj = {"linked": False, "rel": UNKNOWN_REL, "link_text": "—",
                                "match_type": None, "exact": False}

            if proj["linked"]:
                proj_display   = f'✅ {proj["match_type"]}'
                anchor_display = proj["link_text"] if proj["link_text"] not in ("—", "") else ""
                rel = proj["rel"]
                if not rel["known"]:
                    rel_display = "❔ Unknown (not a normal <a> link)"
                else:
                    parts = []
                    if rel["sponsored"]:
                        parts.append("Sponsored ⚠️")
                        reasons.append("Link is rel=sponsored — no SEO value")
                    if rel["nofollow"]:
                        parts.append("Nofollow ⚠️")
                        reasons.append("Link is rel=nofollow — no link juice")
                    if rel["ugc"]:
                        parts.append("UGC")
                        reasons.append("Link is rel=ugc — low SEO value")
                    rel_display = " + ".join(parts) if parts else "Dofollow ✅"
                if not proj["exact"]:
                    proj_display += " (domain only)"
                    reasons.append("Links to the project domain, but not the exact submitted URL")
            elif proj["match_type"] == "short-unresolved":
                proj_display = "⚠️ Short link could not be followed — verify manually"
                manual.append("short link")
                proj_needs_retry = True
                reasons.append(f"Page has a short link ({proj['link_text']}) that couldn't be followed "
                               "automatically — open it to confirm it goes to the project URL")
            elif proj["match_type"] == "text-only":
                proj_display = "⚠️ URL mentioned in text but NOT a clickable link"
                flags.append("Project URL Not a Backlink")
                reasons.append("Project URL is only plain text — no actual <a> link")
            elif is_low_quality_platform(final_url):
                proj_display = "⚠️ Known login-walled platform — verify manually"
                reasons.append("Sngine/PHPDolphin-style platform — outbound links may only show to logged-in users")
                proj_needs_retry = True
            else:
                proj_display = "❌ Not linked"
                flags.append("Project URL Not Linked")
                reasons.append("Backlink missing — page is accessible but does not link to the project URL")
                proj_needs_retry = True

        # ── Content quality + AI likelihood ─────────────────────────────────
        text = _extract_page_text(html)
        if len(text) > 100:
            content = score_content(text)
        if content:
            quality_display = f'{content["score"]}/10 — {content["quality_label"]}'
            ai_display      = f'{content["origin_label"]} ({content["ai_pct"]}% AI)'
            ai_signals      = content["signals"]
            if content["ai_pct"] >= 80:
                flags.append("High AI Content")
                reasons.append(f'Content is {content["ai_pct"]}% likely AI-written')
            if content["score"] < 4:
                reasons.append(f'Low content quality ({content["score"]}/10)')
        else:
            quality_display = ai_display = "— Too short to score"

    # ── Final verdict ───────────────────────────────────────────────────────
    low_q = is_low_quality_platform(final_url)
    if state == "login":
        status = "🚩 Flagged — Login Wall (zero SEO value)"
    elif state == "blocked":
        status = "⚠️ Manual Check"
    elif (low_q or manual) and not flags:
        status = "⚠️ Manual Check"
    elif (low_q or manual) and flags:
        status = "🚩 Flagged + ⚠️ Manual Check"
    else:
        status = "✅ Good" if not flags else "🚩 Flagged"

    result = make_result(
        link_status, status, "; ".join(reasons) if reasons else "All checks passed ✓", method,
        **{
            "Final URL": final_url_out, "Keyword Found?": kw_display,
            "Project URL Linked?": proj_display, "Link Rel": rel_display,
            "Anchor Text": anchor_display, "Content Quality": quality_display,
            "AI Content Score": ai_display, "AI Signals": ai_signals,
        },
    )
    # Ask for a JS-rendered second look if plain HTML missed the link or keyword
    result["_retry"] = state == "live" and fetch["method"] != "browser" and (proj_needs_retry or not kw_found)
    return result


# ══════════════════════════════════════════════════════════════════════════════
# 8. PIPELINE STAGES
# ══════════════════════════════════════════════════════════════════════════════

def http_stage(row: dict):
    """
    Returns ("final", result) or ("browser", reason, fallback_result_or_None).
    """
    live_url, _kw, project_url, activity = _row_fields(row)
    is_onsite = activity.lower() == "on-site blog"

    if not live_url:
        return "final", make_result("⛔ No URL provided", "🚩 Flagged", "No backlink URL provided")
    p = urlparse(live_url)
    if p.scheme not in ("http", "https") or not p.netloc:
        return "final", make_result("❌ Invalid URL format", "🚩 Flagged", f"Invalid URL format: {live_url[:80]}")
    if not is_onsite and project_url and host_of(live_url) == host_of(project_url):
        return "final", make_result("⚠️ Self-link", "🚩 Flagged",
                                    "Self-link — backlink is on the project's own domain (no SEO value)")

    # File submissions — only confirm the link is alive, content is checked by hand
    if activity.lower() in FILE_SUBMISSION_TYPES:
        f = http_fetch(live_url)
        if f["state"] == "dead":
            return "final", make_result(f"❌ {f['error']}", "🚩 Flagged", f"Dead link — {f['error']}", "HTTP")
        status = "✅ Live" if f["state"] == "live" else f"⚠️ {f['error'] or 'Could not confirm'}"
        return "final", make_result(status, "⚠️ Manual Check",
                                    f"File submission ({activity}) — open the file and verify the link by hand",
                                    "HTTP", **{"Keyword Found?": "⚠️ Verify manually (file)",
                                               "Project URL Linked?": "⚠️ Verify manually (file)"})

    if is_js_platform(live_url):
        return "browser", "known JS platform", None

    fetch = http_fetch(live_url)

    if fetch["code"] == 403 and len(fetch["html"]) < 100:
        return "final", make_result(
            "⚠️ IP blocked (HTTP 403)", "⚠️ Manual Check",
            "IP blocked — this site blocks your network entirely, so neither HTTP nor a browser can reach it. "
            "Verify manually in your own browser.", "HTTP")

    reason = needs_browser_reason(fetch, host_of(project_url))
    if reason:
        return "browser", reason, None

    result = analyze(row, fetch)
    if result.pop("_retry", False):
        return "browser", "link/keyword not in raw HTML — rendering JS", result
    return "final", result


def browser_stage(page, row: dict, reason: str, fallback):
    live_url, *_ = _row_fields(row)
    b = browser_fetch(page, live_url)

    if b["state"] == "live":
        challenge = detect_challenge(b["html"])
        if challenge and len(plain_text_fast(b["html"])) < 500:
            if fallback:
                return _with_note(fallback, f"browser hit a CAPTCHA ({challenge}) — using HTTP result")
            return make_result(f"⚠️ CAPTCHA / bot protection ({challenge})", "⚠️ Manual Check",
                               f"CAPTCHA / bot protection detected ({challenge}) — automated tools can't pass it. "
                               "Verify manually in your browser.", "Headless browser")
        result = analyze(row, b, page=page, note=reason)
        result.pop("_retry", None)
        # Still a bot-check page and nothing found → it's the wall, not the page
        wall = looks_like_challenge(b["html"])
        if wall and result["Project URL Linked?"].startswith(("❌", "⛔ N/A (no")):
            challenge = wall
            if fallback:
                return _with_note(fallback, f"browser hit a bot check ({challenge}) — using HTTP result")
            return make_result(f"⚠️ CAPTCHA / bot protection ({challenge})", "⚠️ Manual Check",
                               f"Bot protection detected ({challenge}) — the backlink may exist but automated "
                               "tools can't get past it. Verify manually in your browser.", "Headless browser")
        return result

    if b["state"] in ("login", "dead"):
        result = analyze(row, b, page=None, note=reason)
        result.pop("_retry", None)
        return result

    # Browser couldn't get in either
    if fallback:
        return _with_note(fallback, f"browser failed ({b['error']}) — using HTTP result")
    return last_resort(row, reason, b.get("error"))


def last_resort(row: dict, reason: str, browser_error=None):
    """Alternate HTTP client, then give up with a clear Manual Check."""
    live_url, *_ = _row_fields(row)
    alt = http_fetch(live_url, alt=True)
    if alt["state"] == "live" and len(alt["html"].strip()) > 500 and not detect_challenge(alt["html"]):
        result = analyze(row, alt, note=reason)
        needs_js = result.pop("_retry", False)
        # The browser never rendered this page, so a miss here isn't proof
        if needs_js and result["Audit Result"] in ("🚩 Flagged", "✅ Good"):
            if result["Audit Result"] == "🚩 Flagged":
                result["Audit Result"] = "⚠️ Manual Check"
                result["Flag Reason"] = ("Browser couldn't load this page and the plain-HTML copy may be "
                                         "missing JavaScript content — verify manually. (" + result["Flag Reason"] + ")")
        return result
    if alt["state"] == "dead":
        result = analyze(row, alt, note=reason)
        result.pop("_retry", None)
        return result
    why = browser_error or alt.get("error") or "blocked"
    return make_result(f"⚠️ Verify manually — {why}", "⚠️ Manual Check",
                       f"All automated methods failed ({reason}; {why}). Verify manually in your browser.",
                       "HTTP → browser → alt HTTP")


def _with_note(result: dict, note: str) -> dict:
    result = dict(result)
    result.pop("_retry", None)
    result["Fetch Method"] = f"HTTP — {note}"
    return result


ROW_BUDGET_S = int(os.environ.get("ROW_TIMEOUT", 60))   # hard limit per row in the browser

# Things a backlink check never needs — skipping them makes pages load much faster
_BLOCK_TYPES = {"image", "media", "font"}
_BLOCK_HOSTS = (
    "googletagmanager.com", "google-analytics.com", "doubleclick.net", "googlesyndication.com",
    "adservice.google", "facebook.net", "connect.facebook", "hotjar.com", "clarity.ms",
    "amazon-adsystem.com", "adnxs.com", "taboola.com", "outbrain.com", "criteo.", "quantserve.com",
    "scorecardresearch.com", "pubmatic.com", "rubiconproject.com", "moatads.com", "yandex.ru/metrika",
)


def _route_filter(route):
    try:
        req = route.request
        if req.resource_type in _BLOCK_TYPES or any(h in req.url for h in _BLOCK_HOSTS):
            return route.abort()
        return route.continue_()
    except Exception:
        try:
            route.continue_()
        except Exception:
            pass


def _probe_browser():
    """Try to start headless Chromium once. Returns None if OK, else the error text."""
    if not PLAYWRIGHT_OK:
        return "Playwright is not installed"
    try:
        with sync_playwright() as pw:
            b = pw.chromium.launch(headless=True, args=["--no-sandbox", "--disable-dev-shm-usage"])
            b.close()
        return None
    except Exception as e:
        return str(e)


def ensure_browser() -> bool:
    """
    Make sure Chromium is installed. If it's missing, download it now
    (python -m playwright install chromium) — this happens once per server.
    """
    err = _probe_browser()
    if err is None:
        return True
    if "Executable doesn't exist" not in err:
        return False                       # installed but broken — installing again won't help
    print("   ⏬  Chrome not found — downloading it now (one time, usually under 2 min)...")
    import subprocess
    # --only-shell = just the headless browser (~100 MB instead of ~280 MB)
    for cmd in ([sys.executable, "-m", "playwright", "install", "chromium", "--only-shell"],
                [sys.executable, "-m", "playwright", "install", "chromium"]):
        started, last_print = time.time(), 0
        try:
            proc = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                                    text=True, bufsize=1)
            for line in proc.stdout:
                line = line.strip().replace("\r", " ")
                if "unknown option" in line.lower():
                    break
                # Show download progress every 10s so it never looks frozen
                if line and (time.time() - last_print > 10 or "downloaded" in line.lower()):
                    last_print = time.time()
                    print(f"   ⏬  {line[-110:]}")
                if time.time() - started > 300:
                    print("   ❌  Chrome download took over 5 minutes — giving up for this run.")
                    proc.kill()
                    break
            proc.wait(timeout=30)
        except Exception as e:
            print(f"   ❌  Chrome download failed: {str(e)[:150]}")
        if _probe_browser() is None:
            print(f"   ✅  Chrome installed ({int(time.time() - started)}s).")
            return True
    print("   ❌  Chrome still won't start — rows that need it will be marked Manual Check.")
    return False


def memory_limit_mb():
    """The container's memory limit in MB (Render, Docker), or None on a normal computer."""
    for path in ("/sys/fs/cgroup/memory.max", "/sys/fs/cgroup/memory/memory.limit_in_bytes"):
        try:
            with open(path) as f:
                v = f.read().strip()
            if v != "max" and int(v) < 1 << 50:
                return int(v) >> 20
        except Exception:
            continue
    return None


def safe_browser_workers(requested: int) -> int:
    """Each Chrome worker needs ~350 MB; keep ~600 MB for the web server itself."""
    limit = memory_limit_mb()
    if not limit:
        return requested
    fit = max(1, (limit - 600) // 350)
    if fit < requested:
        print(f"   ℹ️  Server has {limit} MB of memory — using {fit} Chrome worker(s) instead of {requested}")
    if limit < 1000:
        print(f"   ⚠️  {limit} MB is too little for Chrome — it may crash. Use a 2 GB+ server plan.")
    return min(requested, fit)


def _launch_context(pw):
    """Start headless Chrome with the settings every worker uses."""
    launch_args = dict(headless=True, args=[
        "--disable-blink-features=AutomationControlled", "--no-sandbox",
        "--disable-dev-shm-usage", "--disable-features=IsolateOrigins,site-per-process",
        # keep memory down
        "--disable-gpu", "--disable-extensions", "--disable-background-networking",
        "--renderer-process-limit=2", "--js-flags=--max-old-space-size=256",
    ])
    try:
        # Real Google Chrome gets past Cloudflare more often than Chromium
        browser = pw.chromium.launch(channel="chrome", **launch_args)
    except Exception:
        browser = pw.chromium.launch(**launch_args)
    context = browser.new_context(
        user_agent=MAIN_HEADERS["User-Agent"], viewport={"width": 1280, "height": 800},
        locale="en-US", ignore_https_errors=True,
        extra_http_headers={"Accept-Language": "en-US,en;q=0.9", "DNT": "1"},
    )
    context.add_init_script("""
        Object.defineProperty(navigator, 'webdriver', {get: () => undefined});
        Object.defineProperty(navigator, 'plugins', {get: () => [1,2,3,4,5]});
        Object.defineProperty(navigator, 'languages', {get: () => ['en-US','en']});
        window.chrome = { runtime: {} };
    """)
    context.route("**/*", _route_filter)
    context.set_default_timeout(PAGE_TIMEOUT_MS)
    return browser, context


def _short_launch_error(msg: str) -> str:
    if "Executable doesn't exist" in msg:
        return "Chrome is not installed on this server (run: python -m playwright install chromium)"
    if "shared librar" in msg or "missing dependencies" in msg.lower():
        return "Chrome is missing system libraries (deploy with the Dockerfile / Docker runtime)"
    return msg.strip().splitlines()[0][:150] if msg.strip() else "unknown error"


def _browser_child(wid, job_q, res_q):
    """
    One browser worker, running in its own process so the parent can kill it
    if a page freezes Chrome. Messages to the parent:
      ("start", wid, i)  ("done", wid, i, result)  ("launch_error", wid, text)
    """
    try:
        sys.stdout.reconfigure(line_buffering=True)
    except Exception:
        pass
    try:
        print(f"   🔧  Chrome worker {wid}: starting...")
        with sync_playwright() as pw:
            try:
                browser, context = _launch_context(pw)
            except Exception as e:
                res_q.put(("launch_error", wid, _short_launch_error(str(e))))
                return
            page = context.new_page()
            print(f"   🔧  Chrome worker {wid}: ready")
            done_here = 0
            while True:
                job = job_q.get()
                if job is None:
                    break
                i, row, reason, fallback = job
                res_q.put(("start", wid, i))
                print(f"   …  worker {wid} opening row #{i}: {str(row.get('Backlink Live Link', ''))[:70]}")
                # Restart Chrome every 10 rows — long sessions slowly eat memory
                if done_here and done_here % 10 == 0:
                    try:
                        browser.close()
                    except Exception:
                        pass
                    browser, context = _launch_context(pw)
                    page = context.new_page()
                done_here += 1
                try:
                    if page.is_closed():
                        page = context.new_page()
                    result = browser_stage(page, row, reason, fallback)
                except Exception as e:
                    note = str(e).strip().splitlines()[0][:80] if str(e).strip() else "browser error"
                    result = fallback and _with_note(fallback, f"browser: {note} — using HTTP result")
                    result = result or last_resort(row, reason, note)
                res_q.put(("done", wid, i, result))
            browser.close()
    except Exception as e:
        res_q.put(("launch_error", wid, _short_launch_error(str(e))))


def run_browser_phase(jobs, on_done, workers: int):
    """jobs: list of (i, row, reason, fallback). Calls on_done(i, result, label)."""
    if not jobs:
        return
    pending = {j[0] for j in jobs}
    by_index = {j[0]: j for j in jobs}
    launch_errors = []

    def finish(i, result):
        if i in pending:
            pending.discard(i)
            on_done(i, result, "BROWSER")

    if not PLAYWRIGHT_OK:
        launch_errors.append("Playwright is not installed")
    else:
        ensure_browser()
        workers = safe_browser_workers(workers)
        import multiprocessing as mp
        ctx = mp.get_context("spawn")
        job_q, res_q = ctx.Queue(), ctx.Queue()
        for j in jobs:
            job_q.put(j)
        n = max(1, min(workers, len(jobs)))
        for _ in range(n):
            job_q.put(None)                         # one stop signal per worker

        procs, current, dead = {}, {}, set()        # current: wid -> (row index, start time)

        def spawn(wid):
            p = ctx.Process(target=_browser_child, args=(wid, job_q, res_q), daemon=True)
            p.start()
            procs[wid] = p

        for wid in range(1, n + 1):
            spawn(wid)
            time.sleep(0.3)

        last_beat = time.time()
        print(f"   🔧  Starting {n} Chrome worker(s)...")
        while pending:
            if time.time() - last_beat > 30:
                last_beat = time.time()
                busy = ", ".join(f"worker {w}: row #{r} for {int(time.time() - t)}s"
                                 for w, (r, t) in sorted(current.items())) or "no rows open yet"
                alive_n = sum(1 for w, p in procs.items() if p.is_alive())
                print(f"   ⏳  still working — {len(pending)} rows left, {alive_n} worker(s) alive ({busy})")
            try:
                msg = res_q.get(timeout=1)
            except queue.Empty:
                msg = None

            if msg:
                kind, wid = msg[0], msg[1]
                if kind == "start":
                    current[wid] = (msg[2], time.time())
                elif kind == "done":
                    current.pop(wid, None)
                    finish(msg[2], msg[3])
                elif kind == "launch_error":
                    launch_errors.append(msg[2])
                    dead.add(wid)
                    print(f"   ❌  Browser worker {wid} could not start: {msg[2]}")

            now = time.time()
            for wid, p in list(procs.items()):
                if wid in dead:
                    continue
                busy = current.get(wid)
                too_slow = busy and now - busy[1] > ROW_BUDGET_S
                crashed = not p.is_alive() and busy
                if too_slow or crashed:
                    i = busy[0]
                    current.pop(wid, None)
                    if p.is_alive():
                        p.kill()
                    p.join(timeout=5)
                    _i, row, reason, fallback = by_index[i]
                    note = (f"page took over {ROW_BUDGET_S}s — gave up" if too_slow
                            else "browser crashed on this page")
                    result = _with_note(fallback, f"browser: {note} — using HTTP result") if fallback else \
                        make_result(f"⚠️ Verify manually — {note}", "⚠️ Manual Check",
                                    f"Browser could not finish this page ({note}). Verify manually in your browser.",
                                    "Headless browser")
                    finish(i, result)
                    spawn(wid)                      # fresh worker carries on with the queue

            alive = [w for w, p in procs.items() if p.is_alive() and w not in dead]
            if not alive and res_q.empty():
                break                               # every worker finished or failed to start

        for p in procs.values():
            if p.is_alive():
                p.kill()

    # Anything the browser never got to (no Playwright, Chromium missing, crash)
    why = f"headless browser unavailable — {launch_errors[0]}" if launch_errors else "headless browser unavailable"
    if launch_errors and pending:
        print(f"\n   ❌  BROWSER DID NOT START: {launch_errors[0]}")
        print(f"       {len(pending)} rows that need a browser will be marked Manual Check.\n")
    for i, row, reason, fallback in jobs:
        if i in pending:
            result = _with_note(fallback, "no browser available — using HTTP result") if fallback \
                else last_resort(row, reason, why)
            finish(i, result)


# ══════════════════════════════════════════════════════════════════════════════
# 9. EXCEL STYLING
# ══════════════════════════════════════════════════════════════════════════════

def apply_styling(path: str):
    wb = load_workbook(path)
    ws = wb.active
    thin   = Side(style="thin", color="FFCCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fills  = {k: PatternFill("solid", start_color=c) for k, c in
              {"good": CLR_GREEN, "bad": CLR_RED, "warn": CLR_YELLOW, "info": CLR_BLUE}.items()}

    for cell in ws[1]:
        cell.fill      = PatternFill("solid", start_color=CLR_HEADER)
        cell.font      = Font(bold=True, color=CLR_WHITE, size=10, name="Arial")
        cell.border    = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    col_map = {cell.value: cell.column for cell in ws[1]}

    def tick_fill(val):
        return fills["good"] if "✅" in val else fills["warn"] if ("⚠️" in val or "❔" in val) \
            else fills["info"] if ("⛔" in val or "⏸" in val or not val) else fills["bad"]

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True, horizontal="left")
            cell.border    = border
            cell.font      = Font(name="Arial", size=9)

        def cell_of(name):
            idx = col_map.get(name)
            return row[idx - 1] if idx else None

        c = cell_of("Audit Result")
        if c:
            v = str(c.value or "")
            c.fill = fills["good"] if "Good" in v else fills["info"] if "Skipped" in v \
                else fills["bad"] if "Flagged" in v else fills["warn"]
            c.font = Font(bold=True, name="Arial", size=9)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for name in ("Link Status", "Keyword Found?", "Project URL Linked?"):
            c = cell_of(name)
            if c:
                c.fill = tick_fill(str(c.value or ""))

        c = cell_of("Link Rel")
        if c and c.value:
            v = str(c.value)
            c.fill = fills["good"] if "Dofollow" in v else fills["warn"]

        c = cell_of("Content Quality")
        if c:
            v = str(c.value or "")
            c.fill = fills["good"] if ("Excellent" in v or "Good Quality" in v) else fills["warn"] \
                if "Average" in v else fills["bad"] if "Low" in v else fills["info"]

        c = cell_of("AI Content Score")
        if c:
            v = str(c.value or "")
            c.fill = fills["good"] if ("Human" in v) else fills["warn"] if "Mixed" in v \
                else fills["bad"] if ("Likely AI" in v or "AI Generated" in v) else fills["info"]

        for name in ("DA", "SS"):
            c = cell_of(name)
            if c:
                c.alignment = Alignment(horizontal="center", vertical="center")

    widths = {
        "Backlink Create Date": 18, "Website": 16, "Domain Country": 14, "Employee Name": 18,
        "Backlink Activity Type": 20, "Referring Domain": 32, "DA": 10, "SS": 10,
        "Backlink Live Link": 52, "Target Keyword": 28, "Submitted Project Url": 45, "Status": 10,
        "Link Status": 30, "Final URL": 40, "Fetch Method": 28, "Keyword Found?": 24,
        "Project URL Linked?": 30, "Link Rel": 18, "Anchor Text": 26, "Content Quality": 24,
        "AI Content Score": 24, "AI Signals": 50, "Audit Result": 18, "Flag Reason": 60,
        "Email": 26, "User  Name": 14, "Password": 14,
    }
    for name, width in widths.items():
        if name in col_map:
            ws.column_dimensions[get_column_letter(col_map[name])].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


# ══════════════════════════════════════════════════════════════════════════════
# 10. MAIN
# ══════════════════════════════════════════════════════════════════════════════

def main():
    parser = argparse.ArgumentParser(description="Backlink Auditor v4 — Hybrid Engine")
    parser.add_argument("input", help="Excel file, e.g. 'Tijeet singh.xlsx'")
    parser.add_argument("--output", default="", help="Output filename (optional)")
    parser.add_argument("--filter-methods", default="", help="Comma-separated activity types to audit")
    parser.add_argument("--moz-key", default="", help="Moz access_id:secret (reserved for DA/SS)")
    parser.add_argument("--meta-project", default="")
    parser.add_argument("--meta-project-id", default="")
    parser.add_argument("--meta-manager", default="")
    parser.add_argument("--meta-teamlead", default="")
    parser.add_argument("--http-workers", type=int, default=HTTP_WORKERS)
    parser.add_argument("--browser-workers", type=int, default=BROWSER_WORKERS)
    parser.add_argument("--no-browser", action="store_true", help="HTTP only — skip the headless browser")
    if "--install-browser" in sys.argv:                 # used by the web server at startup
        sys.exit(0 if ensure_browser() else 1)
    args = parser.parse_args()

    output = args.output or args.input.replace(".xlsx", "_AUDITED.xlsx")
    filter_methods = [m.strip() for m in args.filter_methods.split(",") if m.strip()]
    use_browser = PLAYWRIGHT_OK and not args.no_browser

    print("\n" + "═" * 64)
    print("   BACKLINK AUDITOR v4 — Hybrid Engine")
    print("═" * 64)
    print(f"   Input   : {args.input}")
    print(f"   Output  : {output}")
    print(f"   Engine  : {args.http_workers} parallel HTTP + "
          + (f"{args.browser_workers} headless Chrome fallback" if use_browser else "no browser (HTTP only)"))
    if not PLAYWRIGHT_OK and not args.no_browser:
        print("   Note    : Playwright not installed — JS pages fall back to HTTP / Manual Check")
    if filter_methods:
        print(f"   Filter  : Auditing only — {', '.join(filter_methods)}")
    print("═" * 64 + "\n")

    xl = pd.ExcelFile(args.input)
    sheet_name = next((s for s in ("Backlinks", "Worksheet") if s in xl.sheet_names), xl.sheet_names[0])
    df = pd.read_excel(args.input, sheet_name=sheet_name, dtype=str).fillna("")
    total = len(df)
    print(f"   Sheet   : {sheet_name}")
    print(f"   📋  {total} backlinks loaded.\n")

    activity_col = next((c for c in df.columns if "activity type" in c.lower() or "method" in c.lower()), None)

    da_prefilled = "DA" in df.columns and df["DA"].str.strip().replace("", pd.NA).notna().any()
    ss_prefilled = "SS" in df.columns and df["SS"].str.strip().replace("", pd.NA).notna().any()
    drop_cols = RESULT_COLUMNS + ["DA (Auto)", "SS % (Auto)", "AI Content Score"]
    if not da_prefilled:
        drop_cols.append("DA")
    if not ss_prefilled:
        drop_cols.append("SS")
    df = df.drop(columns=[c for c in set(drop_cols) if c in df.columns])

    for col in ("Project ID", "Project Manager", "Team Lead"):
        if col not in df.columns:
            pos = df.columns.get_loc("Employee Name") + 1 if "Employee Name" in df.columns else len(df.columns)
            df.insert(pos, col, "")
    for col, val in {"Project ID": args.meta_project_id, "Project Manager": args.meta_manager,
                     "Team Lead": args.meta_teamlead, "Website": args.meta_project}.items():
        if not val:
            continue
        if col in df.columns:
            df[col] = df[col].apply(lambda x: val if str(x).strip() == "" else x)
        else:
            df[col] = val

    results  = [None] * total
    rows     = {}
    to_audit = []
    lock     = threading.Lock()
    done     = [0]

    def report(i, result, label):
        results[i - 1] = result
        ref = str(rows[i].get("Referring Domain", "") or host_of(_row_fields(rows[i])[0])).strip()
        icon = ("✅" if "Good" in result["Audit Result"] else
                "⚠️" if result["Audit Result"].startswith("⚠️") else "🚩")
        with lock:
            done[0] += 1
            print(f"   [{done[0]:>3}/{total}]  {label:<7}  {ref[:38]:<38}  #{i}")
            print(f"          → {icon}  {result['Link Status'][:40]}  |  "
                  f"Proj: {result['Project URL Linked?'][:28]}  |  KW: {result['Keyword Found?'][:20]}")

    for i, (_, row) in enumerate(df.iterrows(), 1):
        rows[i] = row.to_dict()
        method = str(row.get(activity_col, "")).strip() if activity_col else ""
        if filter_methods and method and method not in filter_methods:
            results[i - 1] = make_result(
                f"⏸ Skipped — {method} not selected", "⏸ Skipped",
                f"Activity type '{method}' not selected for this audit run",
                **{k: "⏸ Not audited" for k in ("Keyword Found?", "Project URL Linked?",
                                                 "Content Quality", "AI Content Score")})
            with lock:
                done[0] += 1
                print(f"   [{done[0]:>3}/{total}]  ⏭  Skipped — {method}")
        else:
            to_audit.append(i)

    # ── Phase 1: parallel HTTP ──────────────────────────────────────────────
    print(f"   ⚡ Phase 1 — fast HTTP check of {len(to_audit)} rows ({args.http_workers} at a time)\n")
    browser_jobs = []
    with ThreadPoolExecutor(max_workers=max(1, args.http_workers)) as ex:
        futures = {ex.submit(http_stage, rows[i]): i for i in to_audit}
        for fut in as_completed(futures):
            i = futures[fut]
            try:
                out = fut.result()
            except Exception as e:
                out = ("browser", f"HTTP error: {str(e)[:60]}", None)
            if out[0] == "final":
                report(i, out[1], "HTTP")
            elif use_browser:
                browser_jobs.append((i, rows[i], out[1], out[2]))
            else:
                fb = out[2]
                report(i, _with_note(fb, "browser disabled") if fb else last_resort(rows[i], out[1], "browser disabled"), "HTTP")

    # ── Phase 2: headless browser for the rows that need it ─────────────────
    if browser_jobs:
        browser_jobs.sort(key=lambda j: j[0])
        print(f"\n   🌐 Phase 2 — {len(browser_jobs)} rows need a headless browser "
              f"({args.browser_workers} at a time)\n")
        run_browser_phase(browser_jobs, report, max(1, args.browser_workers))

    results = [r if r is not None else make_result("⚠️ Worker error", "⚠️ Manual Check",
                                                   "Processing error — verify manually") for r in results]
    for r in results:
        r.pop("_retry", None)

    df = pd.concat([df.reset_index(drop=True), pd.DataFrame(results, columns=RESULT_COLUMNS)], axis=1)
    df.to_excel(output, index=False, sheet_name=sheet_name)
    apply_styling(output)

    def count(word):
        return sum(1 for r in results if word in r["Audit Result"])

    print("\n" + "═" * 64)
    print("   AUDIT COMPLETE")
    print("═" * 64)
    print(f"   Total           : {total}")
    print(f"   ✅ Good         : {count('Good')}")
    print(f"   ⚠️  Manual Check : {count('Manual')}")
    print(f"   🚩 Flagged      : {count('Flagged')}")
    if count("Skipped"):
        print(f"   ⏸  Skipped      : {count('Skipped')}")
    print(f"   🌐 Browser used : {len(browser_jobs)} of {len(to_audit)} rows")
    print(f"   📁 Output       : {output}")
    print("═" * 64 + "\n")


if __name__ == "__main__":
    main()
